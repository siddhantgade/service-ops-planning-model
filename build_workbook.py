"""Build Service_Ops_Planning_Model.xlsx with live formulas from the Python model outputs."""
import json
import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule

F = "Arial"
NAVY = "1F3864"
fnt = lambda **k: Font(name=F, size=k.pop("size", 10), **k)
HDR_FILL = PatternFill("solid", fgColor=NAVY)
YEL = PatternFill("solid", fgColor="FFFF00")
GREYF = PatternFill("solid", fgColor="F2F2F2")
thin = Side(style="thin", color="D9D9D9")
BOX = Border(top=thin, bottom=thin, left=thin, right=thin)
BLUE, BLACK, GREEN = "0000FF", "000000", "008000"

fc = pd.read_csv("outputs/forecast_13wk.csv", parse_dates=["week_start"])
bt = pd.read_csv("outputs/backtest_weekly.csv", parse_dates=["week_start"])
acc = pd.read_csv("outputs/accuracy_backtest.csv")
data = pd.read_csv("data/service_ops_demand.csv")
start_fte = float(data["fte_available"].iloc[-4:].mean())
A = dict(pd.read_csv("outputs/assumptions.csv").set_index("assumption")["value"])
ramp_by_week = dict(zip(json.loads(A["change_weeks"]), json.loads(A["change_ramp"])))
peak_weeks = set(json.loads(A["peak_weeks"]))
fte_changes = {int(k): v for k, v in json.loads(A["planned_fte_changes"]).items()}

wb = Workbook()


def header(ws, row, labels, col=1):
    for i, t in enumerate(labels):
        c = ws.cell(row=row, column=col + i, value=t)
        c.font = fnt(bold=True, color="FFFFFF"); c.fill = HDR_FILL; c.border = BOX
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[row].height = 32


def put(ws, ref, value, color=BLACK, fmt=None, bold=False, fill=None, align=None):
    c = ws[ref]; c.value = value; c.font = fnt(color=color, bold=bold); c.border = BOX
    if fmt: c.number_format = fmt
    if fill: c.fill = fill
    if align: c.alignment = Alignment(horizontal=align)
    return c


def widths(ws, ws_widths):
    for i, w in enumerate(ws_widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


# ------------------------------------------------------------------ Guide
g = wb.active; g.title = "Guide"
g["A1"] = "Service Operations Planning Model"; g["A1"].font = fnt(size=14, bold=True, color=NAVY)
lines = [
    ("What this is", "A demand forecast and workforce capacity model for a simulated financial services customer operations team (telephony and back-office cases)."),
    ("Data", "SIMULATED. All volumes, handling times, staffing levels and the provider cap are illustrative assumptions, not real organisational data."),
    ("Forecast", "The 13-week forecast on the Forecast tab comes from the Python model (Holt-Winters). Re-run forecast_and_capacity.py to refresh it."),
    ("Capacity", "The Capacity tab converts forecast demand into workload hours and required FTE for three scenarios, using the Assumptions tab. Everything there is a live formula."),
    ("How to use", "Change the blue cells on the Assumptions tab (yellow = key levers) and watch Capacity and Dashboard update. Pick a scenario (1, 2 or 3) on the Dashboard."),
    ("Colour code", "Blue text = input. Black text = formula. Green text = link from another sheet. Yellow fill = key assumption or selector."),
    ("Scenarios", "1 Base (BAU only). 2 Base plus a planned change project adding cases. 3 Scenario 2 plus a volume uplift and higher absence in peak weeks."),
    ("Method", "Required FTE = workload hours / productive hours per FTE. Productive hours = contracted hours x (1 - shrinkage) x occupancy."),
    ("Gap options", "A shortfall is first covered by overtime (up to the overtime cap per FTE); any remainder is shown as temporary or recruited FTE needed."),
]
for i, (k, v) in enumerate(lines, start=3):
    g.cell(row=i, column=1, value=k).font = fnt(bold=True)
    c = g.cell(row=i, column=2, value=v); c.font = fnt(); c.alignment = Alignment(wrap_text=True, vertical="top")
    g.row_dimensions[i].height = 32
widths(g, [16, 110])

# ------------------------------------------------------------------ Assumptions
a = wb.create_sheet("Assumptions")
a["A1"] = "Assumptions (all illustrative)"; a["A1"].font = fnt(size=12, bold=True, color=NAVY)
header(a, 3, ["Assumption", "Value", "Note"])
params = [
    ("AHT_CALL", "Average handling time per call (min)", float(A["aht_call_min"]), "0.0", "Assumed", True),
    ("AHT_CASE", "Average handling time per back-office case (min)", float(A["aht_case_min"]), "0.0", "Assumed", True),
    ("CONTRACT_HRS", "Contracted hours per FTE per week", float(A["contracted_hours_per_fte"]), "0.0", "Standard 37.5 hour week", False),
    ("SHRINK", "Shrinkage (holidays, sickness, training, meetings)", float(A["shrinkage"]), "0.0%", "Assumed", True),
    ("OCC", "Occupancy (share of productive time on work)", float(A["occupancy"]), "0.0%", "Assumed", True),
    ("OUT_SHARE", "Share of cases routed to outsourced provider", float(A["outsourced_share"]), "0.0%", "Assumed", False),
    ("PROV_CAP", "Provider contracted weekly volume (cases)", float(A["provider_weekly_cap_cases"]), "#,##0", "Assumed contract limit", False),
    ("OT_HRS", "Maximum overtime hours per FTE per week", float(A["overtime_hours_per_fte"]), "0.0", "Assumed", False),
    ("CHANGE_CASES", "Change project: peak extra cases per week", float(A["change_cases_per_week"]), "#,##0", "Scaled by the ramp column below", True),
    ("PEAK_UPLIFT", "Scenario 3: extra volume in peak weeks", float(A["peak_uplift"]), "0.0%", "Stress assumption", True),
    ("PEAK_SHRINK", "Scenario 3: shrinkage in peak weeks", float(A["peak_shrinkage"]), "0.0%", "Winter absence stress assumption", True),
    ("START_FTE", "Starting in-house FTE available", start_fte, "0.0", "Average of last 4 weeks of simulated history", False),
]
for i, (name, label, val, fmt, note, key) in enumerate(params, start=4):
    put(a, f"A{i}", label)
    put(a, f"B{i}", val, color=BLUE, fmt=fmt, fill=YEL if key else None)
    put(a, f"C{i}", note)
    wb.defined_names[name] = DefinedName(name, attr_text=f"Assumptions!$B${i}")
r = 4 + len(params)
put(a, f"A{r}", "Productive hours per FTE per week")
put(a, f"B{r}", "=CONTRACT_HRS*(1-SHRINK)*OCC", fmt="0.00", bold=True)
put(a, f"C{r}", "Formula: contracted hours x (1 - shrinkage) x occupancy")
wb.defined_names["PROD_HRS"] = DefinedName("PROD_HRS", attr_text=f"Assumptions!$B${r}")

PROF = 22
a.cell(row=PROF - 2, column=1, value="Weekly profile for the forecast horizon").font = fnt(bold=True, color=NAVY)
header(a, PROF - 1, ["Week no", "Week start", "Change project ramp", "Peak week (1 = yes)", "Planned FTE change", "Available FTE"])
for w in range(1, 14):
    row = PROF + w - 1
    put(a, f"A{row}", w, align="center")
    put(a, f"B{row}", f"=INDEX(Forecast!$B$2:$B$14,A{row})", color=GREEN, fmt="dd mmm yyyy")
    put(a, f"C{row}", ramp_by_week.get(w, 0.0), color=BLUE, fmt="0.0")
    put(a, f"D{row}", 1 if w in peak_weeks else 0, color=BLUE, align="center")
    put(a, f"E{row}", fte_changes.get(w, 0), color=BLUE, fmt="0;-0;0", align="center")
    put(a, f"F{row}", f"=START_FTE+E{row}" if w == 1 else f"=F{row-1}+E{row}", fmt="0.0")
a.cell(row=PROF + 14, column=1, value="Planned FTE change: assumed leavers (negative) and trained hires (positive). Available FTE is cumulative.").font = fnt(italic=True, color="595959")
widths(a, [52, 14, 40, 18, 18, 14])
a.freeze_panes = "A4"

# ------------------------------------------------------------------ Forecast
f = wb.create_sheet("Forecast")
header(f, 1, ["Week no", "Week start", "Calls forecast", "Calls low (80%)", "Calls high (80%)",
              "Cases BAU forecast", "Cases low (80%)", "Cases high (80%)"])
for i, row in fc.iterrows():
    r = i + 2
    put(f, f"A{r}", int(row.week_no), align="center")
    put(f, f"B{r}", row.week_start.to_pydatetime(), color=BLUE, fmt="dd mmm yyyy")
    for col, key in zip("CDEFGH", ["calls_forecast", "calls_low", "calls_high", "cases_bau_forecast", "cases_bau_low", "cases_bau_high"]):
        put(f, f"{col}{r}", int(row[key]), color=BLUE, fmt="#,##0")
f["A16"] = "Source: Holt-Winters forecast from forecast_and_capacity.py (simulated data). Interval = forecast +/- 1.28 x holdout RMSE. Blue = values pasted from the Python model."
f["A16"].font = fnt(italic=True, color="595959")
widths(f, [10, 14, 16, 16, 16, 18, 16, 16]); f.freeze_panes = "A2"

# ------------------------------------------------------------------ Capacity
c = wb.create_sheet("Capacity")
cols = ["Scenario no", "Scenario", "Week no", "Week start", "Base calls", "Base cases (BAU)", "Change cases",
        "Peak week", "Volume uplift", "Shrinkage", "Calls", "Total cases", "Workload hours",
        "Productive hrs per FTE", "Required FTE", "Available FTE", "Gap FTE", "Overtime hours available",
        "Residual hours after overtime", "Temp or recruit FTE needed", "Provider cases forecast", "Provider over cap (1 = yes)"]
header(c, 1, cols)
c.row_dimensions[1].height = 44
names = {1: "1 Base (BAU only)", 2: "2 Base + change project", 3: "3 Change + peak + absence"}
row = 2
for s in (1, 2, 3):
    for w in range(1, 14):
        r = row
        put(c, f"A{r}", s, align="center"); put(c, f"B{r}", names[s]); put(c, f"C{r}", w, align="center")
        put(c, f"D{r}", f"=INDEX(Forecast!$B$2:$B$14,C{r})", color=GREEN, fmt="dd mmm yyyy")
        put(c, f"E{r}", f"=INDEX(Forecast!$C$2:$C$14,C{r})", color=GREEN, fmt="#,##0")
        put(c, f"F{r}", f"=INDEX(Forecast!$F$2:$F$14,C{r})", color=GREEN, fmt="#,##0")
        put(c, f"G{r}", f"=IF(A{r}>=2,CHANGE_CASES*INDEX(Assumptions!$C${PROF}:$C${PROF+12},C{r}),0)", fmt="#,##0")
        put(c, f"H{r}", f"=IF(AND(A{r}=3,INDEX(Assumptions!$D${PROF}:$D${PROF+12},C{r})=1),1,0)", align="center")
        put(c, f"I{r}", f"=1+H{r}*PEAK_UPLIFT", fmt="0.00")
        put(c, f"J{r}", f"=IF(H{r}=1,PEAK_SHRINK,SHRINK)", fmt="0.0%")
        put(c, f"K{r}", f"=E{r}*I{r}", fmt="#,##0")
        put(c, f"L{r}", f"=(F{r}+G{r})*I{r}", fmt="#,##0")
        put(c, f"M{r}", f"=K{r}*AHT_CALL/60+L{r}*(1-OUT_SHARE)*AHT_CASE/60", fmt="#,##0.0")
        put(c, f"N{r}", f"=CONTRACT_HRS*(1-J{r})*OCC", fmt="0.00")
        put(c, f"O{r}", f"=M{r}/N{r}", fmt="0.0", bold=True)
        put(c, f"P{r}", f"=INDEX(Assumptions!$F${PROF}:$F${PROF+12},C{r})", color=GREEN, fmt="0.0")
        put(c, f"Q{r}", f"=MAX(0,O{r}-P{r})", fmt="0.0")
        put(c, f"R{r}", f"=P{r}*OT_HRS", fmt="#,##0.0")
        put(c, f"S{r}", f"=MAX(0,Q{r}*N{r}-R{r})", fmt="#,##0.0")
        put(c, f"T{r}", f"=IF(S{r}>0,ROUNDUP(S{r}/N{r},0),0)", fmt="0", align="center")
        put(c, f"U{r}", f"=L{r}*OUT_SHARE", fmt="#,##0")
        put(c, f"V{r}", f"=IF(U{r}>PROV_CAP,1,0)", align="center")
        row += 1
c.conditional_formatting.add("Q2:Q40", CellIsRule(operator="greaterThan", formula=["0"], fill=PatternFill("solid", bgColor="F8CBAD")))
c.conditional_formatting.add("V2:V40", CellIsRule(operator="equal", formula=["1"], fill=PatternFill("solid", bgColor="F8CBAD")))
widths(c, [10, 26, 9, 13] + [13] * 18); c.freeze_panes = "E2"

# ------------------------------------------------------------------ Accuracy
k = wb.create_sheet("Accuracy")
k["A1"] = "Forecast accuracy: 13-week holdout (Holt-Winters)"; k["A1"].font = fnt(size=12, bold=True, color=NAVY)
put(k, "A2", "Variance flag threshold (%)"); put(k, "B2", float(A["variance_flag_pct"]), color=BLUE, fmt="0.0", fill=YEL)
header(k, 4, ["Week start", "Calls actual", "Calls forecast", "Calls variance %", "Calls flag",
              "Cases actual", "Cases forecast", "Cases variance %", "Cases flag"])
bc = bt[bt.series == "calls"].reset_index(drop=True); bk = bt[bt.series == "cases_bau"].reset_index(drop=True)
for i in range(len(bc)):
    r = 5 + i
    put(k, f"A{r}", bc.week_start[i].to_pydatetime(), color=BLUE, fmt="dd mmm yyyy")
    put(k, f"B{r}", int(bc.actual[i]), color=BLUE, fmt="#,##0"); put(k, f"C{r}", int(bc.forecast[i]), color=BLUE, fmt="#,##0")
    put(k, f"D{r}", f"=(B{r}-C{r})/B{r}", fmt="0.0%")
    put(k, f"E{r}", f'=IF(ABS(D{r})*100>$B$2,"CHECK","OK")', align="center")
    put(k, f"F{r}", int(bk.actual[i]), color=BLUE, fmt="#,##0"); put(k, f"G{r}", int(bk.forecast[i]), color=BLUE, fmt="#,##0")
    put(k, f"H{r}", f"=(F{r}-G{r})/F{r}", fmt="0.0%")
    put(k, f"I{r}", f'=IF(ABS(H{r})*100>$B$2,"CHECK","OK")', align="center")
last = 5 + len(bc) - 1
put(k, f"A{last+2}", "MAPE", bold=True)
put(k, f"D{last+2}", f"=SUMPRODUCT(ABS(D5:D{last}))/COUNT(D5:D{last})", fmt="0.0%", bold=True)
put(k, f"H{last+2}", f"=SUMPRODUCT(ABS(H5:H{last}))/COUNT(H5:H{last})", fmt="0.0%", bold=True)
put(k, f"A{last+3}", "Bias (forecast vs actual)", bold=True)
put(k, f"D{last+3}", f"=(SUM(C5:C{last})-SUM(B5:B{last}))/SUM(B5:B{last})", fmt="0.0%", bold=True)
put(k, f"H{last+3}", f"=(SUM(G5:G{last})-SUM(F5:F{last}))/SUM(F5:F{last})", fmt="0.0%", bold=True)
k.cell(row=last + 5, column=1, value="Actual and forecast values (blue) come from the Python backtest on simulated data. Seasonal naive MAPE for comparison: calls 4.7%, cases 5.5%.").font = fnt(italic=True, color="595959")
widths(k, [26, 13, 14, 15, 11, 13, 14, 15, 11])

# ------------------------------------------------------------------ Dashboard
d = wb.create_sheet("Dashboard", 1)
d["A1"] = "Dashboard: demand, capacity and scenario view"; d["A1"].font = fnt(size=14, bold=True, color=NAVY)
d["A3"] = "Select scenario (1, 2 or 3)"; d["A3"].font = fnt(bold=True)
put(d, "B3", 2, color=BLUE, fill=YEL, align="center", bold=True)
dv = DataValidation(type="list", formula1='"1,2,3"', allow_blank=False); d.add_data_validation(dv); dv.add("B3")
put(d, "C3", "=INDEX(Capacity!$B$2:$B$40,(B3-1)*13+1)", color=GREEN, bold=True)
d.merge_cells("C3:E3")

d["A5"] = "All scenarios at a glance"; d["A5"].font = fnt(bold=True, color=NAVY)
header(d, 6, ["Scenario", "Peak required FTE", "Weeks short", "Peak gap (FTE)", "Max temp / recruit FTE after overtime", "Weeks provider over cap"])
for i, s in enumerate((1, 2, 3)):
    r = 7 + i
    put(d, f"A{r}", names[s])
    rng = lambda col: f"Capacity!${col}$2:${col}$40"
    put(d, f"B{r}", f"=SUMPRODUCT(MAX(({rng('A')}={s})*{rng('O')}))", fmt="0.0")
    put(d, f"C{r}", f'=COUNTIFS({rng("A")},{s},{rng("Q")},">0")', align="center")
    put(d, f"D{r}", f"=SUMPRODUCT(MAX(({rng('A')}={s})*{rng('Q')}))", fmt="0.0")
    put(d, f"E{r}", f"=SUMPRODUCT(MAX(({rng('A')}={s})*{rng('T')}))", fmt="0", align="center")
    put(d, f"F{r}", f"=SUMIFS({rng('V')},{rng('A')},{s})", fmt="0", align="center")

d["A11"] = "Selected scenario: week by week"; d["A11"].font = fnt(bold=True, color=NAVY)
header(d, 12, ["Week start", "Required FTE", "Available FTE", "Gap FTE", "Temp / recruit FTE needed", "Provider cases", "Provider cap"])
for w in range(1, 14):
    r = 12 + w
    ix = lambda col: f"=INDEX(Capacity!${col}$2:${col}$40,($B$3-1)*13+{w})"
    put(d, f"A{r}", ix("D"), color=GREEN, fmt="dd mmm yyyy")
    put(d, f"B{r}", ix("O"), color=GREEN, fmt="0.0")
    put(d, f"C{r}", ix("P"), color=GREEN, fmt="0.0")
    put(d, f"D{r}", ix("Q"), color=GREEN, fmt="0.0")
    put(d, f"E{r}", ix("T"), color=GREEN, fmt="0", align="center")
    put(d, f"F{r}", ix("U"), color=GREEN, fmt="#,##0")
    put(d, f"G{r}", "=PROV_CAP", fmt="#,##0")
d.conditional_formatting.add("D13:D25", CellIsRule(operator="greaterThan", formula=["0"], fill=PatternFill("solid", bgColor="F8CBAD")))
ch = LineChart(); ch.title = "Required vs available FTE (selected scenario)"; ch.height = 8; ch.width = 18
ch.add_data(Reference(d, min_col=2, max_col=3, min_row=12, max_row=25), titles_from_data=True)
ch.set_categories(Reference(d, min_col=1, min_row=13, max_row=25))
ch.y_axis.title = "FTE"; ch.y_axis.scaling.min = 40; ch.x_axis.number_format = "dd mmm"
ch.series[0].graphicalProperties.line.solidFill = "2A78D6"; ch.series[1].graphicalProperties.line.solidFill = "52514E"
ch.x_axis.delete = False; ch.y_axis.delete = False
d.add_chart(ch, "I3")
widths(d, [28, 16, 14, 14, 22, 16, 14])

wb.save("Service_Ops_Planning_Model.xlsx")
print("saved")
